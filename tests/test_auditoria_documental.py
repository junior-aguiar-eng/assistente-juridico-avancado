"""Auditoria consolidada com documentos sintéticos: dados, conteúdo e pendências."""
import copy
import json
import re
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

import docx
import pytest

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import auditar
import extract
import preencher


def gravar(p, obj):
    p.write_text(json.dumps(obj, ensure_ascii=False, indent=2), encoding="utf-8")
    return p


@pytest.fixture
def pacote(tmp_path):
    fonte = tmp_path / "fonte.txt"
    fonte.write_text("A contratante é Alfa Exemplo Ltda., nesta simulação.", encoding="utf-8")
    contrato = tmp_path / "contrato.docx"
    d = docx.Document()
    d.add_paragraph("Cláusula 1. Confidencialidade por prazo indeterminado.")
    d.save(contrato)
    modelo = tmp_path / "modelo.docx"
    d = docx.Document()
    d.add_paragraph("Contratante: {{contratante}}. Chamado: {{chamado}}.")
    d.save(modelo)
    documentos = [extract.extrair(p)[0] for p in (fonte, contrato, modelo)]
    ext = {"schema_version": "1", "documents": documentos}
    ep = gravar(tmp_path / "extracao.json", ext)
    b = documentos[0]["blocks"][0]
    campos = {"campos": [
        {"nome": "contratante", "estado": "extraido", "valor": "Alfa Exemplo Ltda.",
         "fonte": {"document_id": documentos[0]["document_id"], "locator": b["locator"], "trecho": b["text"]}},
        {"nome": "chamado", "estado": "pendente", "motivo": "Sem comprovante do chamado."}]}
    cp = gravar(tmp_path / "campos.json", campos)
    minuta = tmp_path / "minuta.docx"
    preencher.gerar(campos, modelo, minuta, ext)
    b = documentos[1]["blocks"][0]
    issues = {"schema_version": "1", "posicao_da_parte": "contratante", "issues": [
        {"id": "i1", "titulo": "Preferência de duração", "descricao": "Negociar a duração conforme instrução comercial.",
         "clausula": "1", "clause_text": b["text"], "document_id": documentos[1]["document_id"], "locator": b["locator"],
         "severidade_juridica": "nao_avaliada", "motivo_nao_avaliada": "Não houve avaliação jurídica.",
         "severidade_negocial": "media", "base": [{"tipo": "pratica_do_cliente", "nota": "Preferência fictícia de 24 meses."}]}]}
    ip = gravar(tmp_path / "issues.json", issues)
    fp = gravar(tmp_path / "facts.json", {"schema_version": "1", "facts": []})
    return {"base": tmp_path, "ext": ep, "facts": fp, "issues": ip, "campos": cp, "modelo": modelo, "minuta": minuta}


def executar(p, **kw):
    args = {k: p[k] for k in ("issues", "campos", "modelo", "minuta")}
    args.update(kw)
    return auditar.auditar(p["facts"], None, [p["ext"]], base=p["base"], **args)


def gates(q):
    return {g["id"]: g["resultado"] for g in q["gates"]}


def test_relatorio_real_de_redline_com_issues_nao_e_outro_conjunto(pacote):
    import redline_docx
    ips = json.loads(pacote['issues'].read_text(encoding='utf-8'))
    original = pacote['base']/'contrato.docx'
    saida = pacote['base']/'contrato-redline.docx'
    edicoes = [{'acao':'substituir','de':ips['issues'][0]['clause_text'],
                'para':'Cláusula 1. Prazo sujeito a negociação.','issue_id':'i1'}]
    r = redline_docx.aplicar(original,edicoes,saida,issues=pacote['issues'])
    assert r['aplicadas'] == 1 and r['recusadas'] == 0
    rp = gravar(pacote['base']/'relatorio-redline.json',r)
    q = executar(pacote,extras=[rp])
    assert gates(q)['G14_apontamentos_contratuais'] == 'alerta'
    assert q['documentos']['apontamentos']['entregavel'] is False
    assert 'relatorio-redline.json' in {a['caminho'] for a in q['artefatos']}


def test_issues_malformado_com_string_continua_reprovado(pacote):
    ip = gravar(pacote['base']/'malformado.json',{'issues':'arquivo.json'})
    q = executar(pacote,issues=ip)
    assert gates(q)['G14_apontamentos_contratuais'] == 'reprovado'


def test_consolida_validadores_e_preserva_pendencias_sem_aprovar(pacote):
    q = executar(pacote)
    assert q["estado_tecnico"] == "revisao_pronta" and q["faixa"] == "mista"
    assert gates(q)["G14_apontamentos_contratuais"] == "alerta"
    assert gates(q)["G15_minuta"] == "alerta"
    assert q["documentos"]["apontamentos"]["avaliacao_juridica_pendente"] == ["i1"]
    assert q["documentos"]["minuta"]["completo"] is False
    assert q["documentos"]["minuta"]["marcados_no_documento"] == {"chamado": "pendente"}
    assert q["aprovacao_humana_obrigatoria"] is True
    assert {p["item"] for p in q["pendencias"]} >= {"i1", "chamado"}
    paths = {a["caminho"] for a in q["artefatos"]}
    assert {"issues.json", "campos.json", "modelo.docx", "minuta.docx"} <= paths
    qp = gravar(pacote["base"] / "qc.json", q)
    s = auditar.status(qp)
    assert s["documentos"] == q["documentos"]
    assert s["pendencias"] == q["pendencias"] and not s["liberado_para_uso"]


def test_apontamento_comercial_com_baixo_risco_juridico_bloqueia(pacote):
    obj = json.loads(pacote["issues"].read_text(encoding="utf-8"))
    obj["issues"][0]["severidade_juridica"] = "baixa"
    gravar(pacote["issues"], obj)
    q = executar(pacote)
    assert q["estado_tecnico"] == "bloqueado"
    assert gates(q)["G14_apontamentos_contratuais"] == "reprovado"


def test_valor_sem_trecho_no_preenchimento_bloqueia(pacote):
    obj = json.loads(pacote["campos"].read_text(encoding="utf-8"))
    obj["campos"][0]["valor"] = "Empresa que não está na fonte"
    gravar(pacote["campos"], obj)
    q = executar(pacote)
    assert q["estado_tecnico"] == "bloqueado" and gates(q)["G15_minuta"] == "reprovado"


def test_minuta_que_nao_corresponde_ao_modelo_e_campos_bloqueia(pacote):
    d = docx.Document(pacote["minuta"])
    d.add_paragraph("Texto acrescentado sem relação com o modelo ou os dados.")
    d.save(pacote["minuta"])
    q = executar(pacote)
    assert gates(q)["G15_minuta"] == "reprovado"
    assert q["estado_tecnico"] == "bloqueado"


@pytest.mark.parametrize("faltante", ["modelo", "minuta"])
def test_campos_exigem_modelo_e_minuta_para_conferencia_final(pacote, faltante):
    q = executar(pacote, **{faltante: None})
    assert gates(q)["G15_minuta"] == "reprovado"


@pytest.mark.parametrize("nome", ["issues", "campos"])
@pytest.mark.parametrize("conteudo", ["{invalido", "[]", '{"outro": []}'])
def test_arquivo_invalido_gera_gate_em_vez_de_sucesso_ou_excecao(pacote, nome, conteudo):
    pacote[nome].write_text(conteudo, encoding="utf-8")
    q = executar(pacote)
    assert q["estado_tecnico"] == "bloqueado"


def test_reconhece_json_contratual_e_campos_em_artefatos_com_nome_arbitrario(pacote):
    ip = pacote["issues"].rename(pacote["base"] / "avaliacao.json")
    cp = pacote["campos"].rename(pacote["base"] / "dados.json")
    q = executar(pacote, issues=None, campos=None, extras=[ip, cp])
    assert gates(q)["G14_apontamentos_contratuais"] == "alerta"
    assert gates(q)["G15_minuta"] == "alerta"


def test_dois_jsons_de_campos_nao_sao_selecionados_silenciosamente(pacote):
    extra = pacote["base"] / "outros-campos.json"
    extra.write_bytes(pacote["campos"].read_bytes())
    with pytest.raises(ValueError, match="[Aa]mb|[Mm]ais de um"):
        executar(pacote, extras=[extra])


def test_campos_completos_e_minuta_correspondente_passam_sem_pendencia_de_campo(pacote):
    obj = json.loads(pacote["campos"].read_text(encoding="utf-8"))
    obj["campos"][1] = {"nome": "chamado", "estado": "informado", "valor": "CH-001", "informado_em": "2026-10-20", "fala": "O chamado é CH-001."}
    gravar(pacote["campos"], obj)
    preencher.gerar(obj, pacote["modelo"], pacote["minuta"], json.loads(pacote["ext"].read_text(encoding="utf-8")))
    q = executar(pacote)
    assert gates(q)["G15_minuta"] == "aprovado"
    assert q["documentos"]["minuta"]["completo"] is True
    assert q["documentos"]["minuta"]["marcados_no_documento"] == {}


def test_status_detecta_alteracao_posterior_do_modelo(pacote):
    q = executar(pacote)
    qp = gravar(pacote["base"] / "qc.json", q)
    with pacote["modelo"].open("ab") as f:
        f.write(b"alterado depois da auditoria")
    assert "modelo.docx" in auditar.status(qp)["alterados"]


def test_cli_consolida_documentos_e_status(pacote):
    qp = pacote["base"] / "qc-cli.json"
    cmd = [sys.executable, "-B", str(RAIZ / "scripts/auditar.py"), "auditar", "--facts", str(pacote["facts"]),
           "--extracoes", str(pacote["ext"]), "--issues", str(pacote["issues"]), "--campos", str(pacote["campos"]),
           "--modelo", str(pacote["modelo"]), "--minuta", str(pacote["minuta"]), "--tipo-entrega", "interna", "--saida", str(qp)]
    p = subprocess.run(cmd, cwd=pacote["base"], capture_output=True)
    assert p.returncode == 0, p.stderr.decode(errors="replace")
    assert json.loads(p.stdout.decode("utf-8"))["documentos"]["minuta"]["completo"] is False


def test_conteudo_fora_do_corpo_tambem_precisa_corresponder_ao_modelo(pacote):
    with ZipFile(pacote["minuta"]) as z:
        partes = {n: z.read(n) for n in z.namelist()}
    antes = partes["word/styles.xml"]
    partes["word/styles.xml"] = re.sub(br'<w:sz w:val="[0-9]+"/>', b'<w:sz w:val="72"/>', antes, count=1)
    assert partes["word/styles.xml"] != antes
    with ZipFile(pacote["minuta"], "w") as z:
        for n, dados in partes.items():
            z.writestr(n, dados)
    q = executar(pacote)
    assert gates(q)["G15_minuta"] == "reprovado"


@pytest.mark.parametrize("estado,dados", [
    ("conflito", {"candidatos": [{"valor": "A"}, {"valor": "B"}]}),
    ("inferido", {"valor": "CH-001", "razao": "Hipótese sem prova, não deve ser aceita."}),
    ("nao_se_aplica", {"motivo": "Requer decisão sobre o destino do campo."})])
def test_estados_marcados_nao_se_tornam_minuta_completa(pacote, estado, dados):
    obj = json.loads(pacote["campos"].read_text(encoding="utf-8"))
    obj["campos"][1] = {"nome": "chamado", "estado": estado, **dados}
    gravar(pacote["campos"], obj)
    preencher.gerar(obj, pacote["modelo"], pacote["minuta"], json.loads(pacote["ext"].read_text(encoding="utf-8")))
    q = executar(pacote)
    assert gates(q)["G15_minuta"] == "alerta"
    assert q["documentos"]["minuta"]["completo"] is False
    assert q["documentos"]["minuta"]["marcados_no_documento"] == {"chamado": estado}


def test_juntar_extracoes_separadas_preserva_conferencia_dos_campos_e_clausulas(pacote):
    obj = json.loads(pacote["ext"].read_text(encoding="utf-8"))
    eps = [gravar(pacote["base"] / f"extracao-{i}.json", {"documents": [d]}) for i, d in enumerate(obj["documents"])]
    q = auditar.auditar(pacote["facts"], None, eps, base=pacote["base"],
                       **{k: pacote[k] for k in ("issues", "campos", "modelo", "minuta")})
    assert gates(q)["G14_apontamentos_contratuais"] == "alerta" and gates(q)["G15_minuta"] == "alerta"


def test_regra_inativa_nao_e_contornada_por_nao_avaliada(pacote):
    obj = json.loads(pacote["issues"].read_text(encoding="utf-8"))
    obj["issues"][0]["base"] = [{"tipo": "regra", "rule_id": "civ-422"}]
    gravar(pacote["issues"], obj)
    q = executar(pacote)
    assert gates(q)["G14_apontamentos_contratuais"] == "reprovado"
    assert "regra_nao_revisada" in {e["codigo"] for e in q["documentos"]["apontamentos"]["erros"]}


@pytest.mark.parametrize("nome", ["issues", "campos", "modelo", "minuta"])
def test_entrada_documental_ausente_bloqueia_com_gate(pacote, nome):
    q = executar(pacote, **{nome: pacote["base"] / "ausente"})
    assert q["estado_tecnico"] == "bloqueado"


def test_pacotes_de_regras_entram_no_manifesto_de_integridade(pacote):
    q = executar(pacote)
    assert {Path(a["caminho"]).name for a in q["artefatos"] if a["tipo"] == "regras"} == {"servicos_saas.json", "lgpd.json"}


@pytest.mark.parametrize("chave", ["campos", "issues"])
@pytest.mark.parametrize("valor", ["invalido", 7, None, {}])
def test_json_documental_mal_tipado_nao_escapa_como_artefato_generico(pacote, chave, valor):
    p = gravar(pacote["base"] / "malformado.json", {chave: valor})
    q = executar(pacote, issues=None, campos=None, modelo=None, minuta=None, extras=[p])
    assert q["estado_tecnico"] == "bloqueado"
    gid = "G15_minuta" if chave == "campos" else "G14_apontamentos_contratuais"
    assert gates(q)[gid] == "reprovado"


def test_resultado_do_validador_com_contagem_issues_nao_e_entrada_documental(pacote):
    rel = gravar(pacote["base"] / "validacao.json", {"erros": [], "avisos": [], "issues": 2, "rascunho": False, "entregavel": False})
    q = executar(pacote, extras=[rel])
    assert q["estado_tecnico"] == "revisao_pronta"
