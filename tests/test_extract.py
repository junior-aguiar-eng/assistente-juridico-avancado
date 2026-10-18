import json
import subprocess
import sys
from pathlib import Path

import pytest

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import extract  # noqa: E402

CASOS = RAIZ / "tests" / "casos"
PDF_STJ = CASOS / "stj_aresp_2432194.pdf"
EMENTA_REF = CASOS / "stj_aresp_2432194_ementa_ref.txt"


@pytest.fixture(scope="module")
def stj():
    doc, _ = extract.extrair(PDF_STJ)
    return doc


def _bloco(doc, trecho):
    achados = [b for b in doc["blocks"] if trecho in b["text"]]
    assert achados, f"trecho não encontrado: {trecho!r}"
    return achados[0]


# ---------------------------------------------------------------- PDF real do STJ
def test_stj_documento_e_status(stj):
    assert stj["status"] == "ok"
    assert stj["page_count"] == 16
    assert stj["document_id"] == "doc-" + stj["sha256"][:12]
    assert stj["idioma_assumido"] == "pt-BR"


def test_stj_secoes_na_ordem_e_paginas(stj):
    seg = extract.segmentos(stj)
    assert [s["section"] for s in seg] == [
        "cabecalho", "ementa", "acordao", "ementa", "relatorio", "voto", "termo_de_julgamento", "rodape"]
    por = {s["section"]: s for s in seg if s["section"] in ("relatorio", "voto", "termo_de_julgamento")}
    assert por["relatorio"]["page"] == 4
    assert por["voto"]["page"] == 5 and por["voto"]["page_end"] == 14
    assert por["termo_de_julgamento"]["page"] == 15


def test_stj_ementa_extraida_confere_com_fonte_estruturada(stj):
    ref = EMENTA_REF.read_text(encoding="utf-8")
    ementas = [s for s in extract.segmentos(stj) if s["section"] == "ementa"]
    assert len(ementas) == 2  # o PDF reúne dois documentos, cada um com a ementa
    for s in ementas:
        assert extract.similaridade(s["text"], ref) >= 0.99


def test_stj_rodape_sai_do_corpo_e_volta_como_blocos_da_secao_rodape(stj):
    corpo = " ".join(b["text"] for b in stj["blocks"] if b["section"] != "rodape")
    assert "Documento eletrônico" not in corpo and "Código de Controle" not in corpo
    rodape = [b for b in stj["blocks"] if b["section"] == "rodape"]
    assert rodape and all(b["tags"] == ["repetido"] and b["locator"].startswith("p.") for b in rodape)
    publicacao = next(b for b in rodape if "Publicação no DJEN/CNJ de 20/03/2026" in b["text"])
    assert publicacao["locator"] == "p.1-2"  # a data de publicação continua citável
    assert any(a["codigo"] == "cabecalho_rodape_removido" and "rodape" in a["detalhe"] for a in stj["warnings"])
    assert "layout_em_colunas_ou_tabela" not in {a["codigo"] for a in stj["warnings"]}


def test_stj_citacao_distinta_do_texto_do_relator(stj):
    corte_local = _bloco(stj, "Na espécie, Petromais")
    assert corte_local["is_quote"] and corte_local["quote_basis"] == "recuo"
    relator = _bloco(stj, "JOSÉ alegou ofensa ao art. 205 do CC/2002. Sustentou")
    assert not relator["is_quote"]
    introducao = _bloco(stj, "Sobre o tema, a Corte local consignou")
    assert not introducao["is_quote"]


def test_stj_ementas_transcritas_de_precedentes_sao_citacao(stj):
    voto = [b for b in stj["blocks"] if b["section"] == "voto"]
    transcritos = [b for b in voto if "AREsp n. 2.293.400/SP" in b["text"] or "AgInt no REsp n. 1.878.110" in b["text"]]
    assert transcritos and all(b["is_quote"] for b in transcritos)
    # mas o texto próprio do relator no voto não é marcado como citação
    proprios = [b for b in voto if b["text"].startswith(("Nesse contexto", "Desse modo", "Portanto, não foram"))]
    assert len(proprios) == 3 and not any(b["is_quote"] for b in proprios)


def test_stj_art_206_aparece_em_citacao_e_fora_dela(stj):
    voto = [b for b in stj["blocks"] if b["section"] == "voto" and "206" in b["text"]]
    assert any(b["is_quote"] for b in voto) and any(not b["is_quote"] for b in voto)


def test_stj_dispositivo_marcado_com_pagina(stj):
    disp = [b for b in stj["blocks"] if "dispositivo" in b["tags"]]
    assert len(disp) == 1
    assert disp[0]["locator"] == "p.14" and "NEGAR-LHE PROVIMENTO" in disp[0]["text"]


def test_stj_blocos_tem_campos_obrigatorios(stj):
    ids = [b["id"] for b in stj["blocks"]]
    assert len(ids) == len(set(ids))
    for b in stj["blocks"]:
        assert b["text"].strip() and b["locator"].startswith("p.") and b["section"]
        assert isinstance(b["is_quote"], bool) and isinstance(b["tags"], list)


def test_stj_sem_caracteres_de_substituicao(stj):
    assert not any("�" in b["text"] for b in stj["blocks"])


# ---------------------------------------------------------------- CLI e cache
def _cli(*args):
    r = subprocess.run([sys.executable, "-P", str(RAIZ / "scripts" / "extract.py"), *map(str, args)],
                       capture_output=True, cwd=RAIZ)
    return r.returncode, json.loads(r.stdout.decode("utf-8")), r.stderr.decode("utf-8", "replace")


def test_cli_stdout_somente_json_e_cache(tmp_path):
    cache = tmp_path / "cache"
    cod1, j1, _ = _cli(PDF_STJ, "--cache-dir", cache)
    assert cod1 == 0 and j1["run"]["cache_hits"] == 0 and j1["documents"][0]["status"] == "ok"
    cod2, j2, _ = _cli(PDF_STJ, "--cache-dir", cache)
    assert j2["run"]["cache_hits"] == 1
    assert j2["documents"][0]["blocks"] == j1["documents"][0]["blocks"]


def test_cache_invalida_quando_um_byte_muda(tmp_path):
    copia = tmp_path / "copia.pdf"
    copia.write_bytes(PDF_STJ.read_bytes())
    cache = tmp_path / "cache"
    d1, hit1 = extract.extrair(copia, cache)
    copia.write_bytes(PDF_STJ.read_bytes() + b"\n")
    d2, hit2 = extract.extrair(copia, cache)
    assert (hit1, hit2) == (False, False)
    assert d1["document_id"] != d2["document_id"]
    assert d2["status"] == "ok" and len(list(cache.glob("*.json"))) == 2


def test_cli_codigo_de_saida_2_quando_ha_falha(tmp_path):
    ruim = tmp_path / "x.pdf"
    ruim.write_bytes(b"%PDF-1.4 isto nao e um pdf")
    cod, j, _ = _cli(ruim, "--no-cache")
    assert cod == 2 and j["documents"][0]["status"] == "failed"


# ---------------------------------------------------------------- falhas explícitas
def test_pdf_sem_texto_e_no_text(tmp_path):
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(595, 842)
    w.add_blank_page(595, 842)
    f = tmp_path / "escaneado.pdf"
    with open(f, "wb") as fh:
        w.write(fh)
    doc, _ = extract.extrair(f)
    assert doc["status"] == "no_text" and doc["blocks"] == [] and doc["page_count"] == 2
    assert [a["pagina"] for a in doc["warnings"] if a["codigo"] == "pagina_sem_texto"] == [1, 2]


def test_pdf_corrompido_e_truncado_nao_levantam_excecao(tmp_path):
    f = tmp_path / "corrompido.pdf"
    f.write_bytes(b"%PDF-1.4\n1 0 obj << /Type /Catalog >> endobj\n")
    doc, _ = extract.extrair(f)
    assert doc["status"] == "failed" and doc["warnings"][0]["codigo"] == "erro_de_leitura"
    assert "No /Root object" in doc["warnings"][0]["detalhe"]
    for corte in (20000, 60000, len(PDF_STJ.read_bytes()) - 200):  # truncado em qualquer ponto: falha explícita, nunca texto parcial
        trunc = tmp_path / f"truncado_{corte}.pdf"
        trunc.write_bytes(PDF_STJ.read_bytes()[:corte])
        doc, _ = extract.extrair(trunc)
        assert doc["status"] == "failed" and doc["blocks"] == [] and doc["warnings"][0]["codigo"] == "erro_de_leitura"


def test_entradas_recusadas(tmp_path):
    exe = tmp_path / "a.exe"
    exe.write_bytes(b"MZ")
    assert extract.extrair(exe)[0]["warnings"][0]["codigo"] == "extensao_nao_aceita"
    assert extract.extrair(tmp_path / "nao_existe.pdf")[0]["warnings"][0]["codigo"] == "arquivo_inexistente"
    grande = tmp_path / "g.txt"
    grande.write_bytes(b"a" * 2048)
    assert extract.extrair(grande, max_mb=0.001)[0]["warnings"][0]["codigo"] == "arquivo_grande"


# ---------------------------------------------------------------- DOCX, TXT, MD
def test_docx_paragrafos_tabela_citacao_e_alteracoes(tmp_path):
    import docx
    from docx.shared import Cm

    d = docx.Document()
    d.add_heading("Cláusula Primeira", level=1)
    d.add_paragraph("O contrato tem prazo de 12 meses.")
    q = d.add_paragraph("Texto transcrito de outro documento.")
    q.paragraph_format.left_indent = Cm(2.5)
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Valor"
    t.cell(1, 0).text, t.cell(1, 1).text = "Licença", "R$ 1.000,00"
    f = tmp_path / "c.docx"
    d.save(f)
    doc, _ = extract.extrair(f)
    assert doc["status"] == "ok" and doc["page_count"] is None
    locs = {b["locator"]: b for b in doc["blocks"]}
    assert locs["par.2"]["is_quote"] is False and locs["par.2"]["section"] == "clausula_primeira"
    assert locs["par.3"]["is_quote"] is True
    assert locs["tab.1.lin.2"]["text"] == "Licença | R$ 1.000,00"
    assert not doc["warnings"]


def test_docx_com_alteracoes_rastreadas_gera_aviso(tmp_path):
    import docx
    import zipfile

    d = docx.Document()
    d.add_paragraph("Texto.")
    base = tmp_path / "b.docx"
    d.save(base)
    out = tmp_path / "rastreado.docx"
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                data = data.replace(b"</w:body>", b'<w:ins w:id="1" w:author="x" w:date="2026-01-01T00:00:00Z"/></w:body>')
            zout.writestr(item, data)
    doc, _ = extract.extrair(out)
    assert any(a["codigo"] == "alteracoes_rastreadas" for a in doc["warnings"])


def test_txt_e_md(tmp_path):
    md = tmp_path / "n.md"
    md.write_text("# Fatos\n\nO réu foi citado.\n\n> citação do acórdão\n> segunda linha\n", encoding="utf-8")
    doc, _ = extract.extrair(md)
    por = {b["locator"]: b for b in doc["blocks"]}
    assert por["par.2"]["section"] == "fatos" and not por["par.2"]["is_quote"]
    assert por["par.3"]["is_quote"] and por["par.3"]["quote_basis"] == "marcador"
    latin = tmp_path / "l.txt"
    latin.write_bytes("Ação de cobrança".encode("windows-1252"))
    doc, _ = extract.extrair(latin)
    assert doc["blocks"][0]["text"] == "Ação de cobrança"
    assert doc["warnings"][0]["codigo"] == "codificacao_assumida"
    vazio = tmp_path / "v.txt"
    vazio.write_text("  \n\n", encoding="utf-8")
    assert extract.extrair(vazio)[0]["status"] == "no_text"


def test_similaridade_ignora_caixa_e_pontuacao_mas_detecta_troca():
    assert extract.similaridade("Art. 206, § 5º, I.", "ART 206 5º I") == 1.0
    assert extract.similaridade("prazo de cinco anos", "prazo de dez anos") < 0.9


# ---------------------------------------------------------------- correções da revisão crítica (A1, C3)
def test_text_sha256_nao_depende_da_segmentacao_em_blocos(monkeypatch):
    normal, _ = extract.extrair(PDF_STJ)
    monkeypatch.setattr(extract, "_salto_de_paragrafo", lambda *a, **k: True)
    fatiado, _ = extract.extrair(PDF_STJ)
    assert len(fatiado["blocks"]) > len(normal["blocks"])  # segmentação diferente
    assert fatiado["text_sha256"] == normal["text_sha256"]


def _docx_modificado(tmp_path, nome, inserir_no_corpo="", partes_extras=None, cabecalho=None):
    import docx
    import zipfile

    d = docx.Document()
    d.add_paragraph("Texto visível.")
    if cabecalho:
        d.sections[0].header.paragraphs[0].text = cabecalho
    base = tmp_path / f"{nome}_base.docx"
    d.save(base)
    out = tmp_path / f"{nome}.docx"
    with zipfile.ZipFile(base) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            dados = zin.read(item.filename)
            if item.filename == "word/document.xml":
                dados = dados.replace(b"</w:body>", inserir_no_corpo.encode("utf-8") + b"</w:body>")
            zout.writestr(item, dados)
        for nome_parte, conteudo in (partes_extras or {}).items():
            zout.writestr(nome_parte, conteudo)
    return out


W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def test_docx_alteracao_rastreada_le_versao_aceita_e_vira_partial(tmp_path):
    corpo = ('<w:p><w:ins w:id="1" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:t>Inserção importante.</w:t></w:r></w:ins></w:p>'
             '<w:p><w:r><w:t>Mantido </w:t></w:r><w:del w:id="2" w:author="a" w:date="2026-01-01T00:00:00Z"><w:r><w:delText>apagado</w:delText></w:r></w:del></w:p>')
    f = _docx_modificado(tmp_path, "rastreado", corpo)
    doc, _ = extract.extrair(f)
    assert doc["status"] == "partial"
    textos = {b["text"]: b for b in doc["blocks"]}
    assert "Inserção importante." in textos and "inserido" in textos["Inserção importante."]["tags"]
    assert "Mantido" in textos and "apagado" not in " ".join(textos)
    av = next(a for a in doc["warnings"] if a["codigo"] == "alteracoes_rastreadas")
    assert "1 inserção" in av["detalhe"] and "1 exclusão" in av["detalhe"]


def test_docx_cabecalho_e_caixa_de_texto_nao_extraidos_geram_aviso_e_partial(tmp_path):
    caixa = '<w:p><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>Texto da caixa</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p>'
    f = _docx_modificado(tmp_path, "partes", caixa, cabecalho="Processo 123 - CABEÇALHO")
    doc, _ = extract.extrair(f)
    assert doc["status"] == "partial"
    assert not any("caixa" in b["text"].lower() for b in doc["blocks"])
    detalhes = " ".join(a["detalhe"] for a in doc["warnings"] if a["codigo"] == "parte_nao_extraida")
    assert "caixas de texto" in detalhes and "CABEÇALHO" in detalhes


def test_docx_notas_de_rodape_e_controle_de_conteudo_sao_extraidos(tmp_path):
    sdt = '<w:sdt><w:sdtContent><w:p><w:r><w:t>Texto em controle de conteúdo.</w:t></w:r></w:p></w:sdtContent></w:sdt>'
    notas = (f'<?xml version="1.0" encoding="UTF-8"?><w:footnotes {W}>'
             '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>'
             '<w:footnote w:id="2"><w:p><w:r><w:t>Nota: ver fls. 12.</w:t></w:r></w:p></w:footnote></w:footnotes>')
    f = _docx_modificado(tmp_path, "notas", sdt, {"word/footnotes.xml": notas})
    doc, _ = extract.extrair(f)
    assert doc["status"] == "ok" and not doc["warnings"]
    por_loc = {b["locator"]: b for b in doc["blocks"]}
    assert por_loc["nota.2"]["text"] == "Nota: ver fls. 12." and por_loc["nota.2"]["section"] == "notas_de_rodape"
    assert any(b["text"] == "Texto em controle de conteúdo." for b in doc["blocks"])


def test_docx_celula_mesclada_nao_duplica_texto(tmp_path):
    import docx

    d = docx.Document()
    t = d.add_table(rows=1, cols=3)
    t.cell(0, 0).merge(t.cell(0, 1)).text = "Mesclada"
    t.cell(0, 2).text = "Valor"
    f = tmp_path / "m.docx"
    d.save(f)
    doc, _ = extract.extrair(f)
    assert [b["text"] for b in doc["blocks"]] == ["Mesclada | Valor"]


def test_docx_com_tamanho_descomprimido_excessivo_e_recusado(tmp_path, monkeypatch):
    f = _docx_modificado(tmp_path, "grande")
    monkeypatch.setattr(extract, "MAX_DESCOMPRIMIDO", 1000)
    doc, _ = extract.extrair(f)
    assert doc["status"] == "failed" and doc["warnings"][0]["codigo"] == "arquivo_suspeito"


def _pdf_cru(caminho, paginas):
    """PDF mínimo; cada página é ('imagem', texto) ou ('texto', texto). Sem dependências."""
    objs, kids = [], []
    # 1 catálogo, 2 páginas, 3 fonte, 4 imagem; depois (página, conteúdo) por página
    fixos = [b"<< /Type /Catalog /Pages 2 0 R >>", None, b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
             b"<< /Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 12 >>\nstream\n"
             + bytes(range(12)) + b"\nendstream"]
    objs = list(fixos)
    for i, (tipo, texto) in enumerate(paginas):
        n_pag = 5 + 2 * i
        kids.append(f"{n_pag} 0 R")
        desenho = "q 535 0 0 780 30 30 cm /Im0 Do Q " if tipo == "imagem" else ""
        y = 20 if tipo == "imagem" else 700
        cont = f"{desenho}BT /F1 9 Tf 72 {y} Td ({texto}) Tj ET".encode("latin-1")
        objs.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] /Contents {n_pag + 1} 0 R "
                    f"/Resources << /Font << /F1 3 0 R >> /XObject << /Im0 4 0 R >> >> >>".encode())
        objs.append(b"<< /Length %d >>\nstream\n" % len(cont) + cont + b"\nendstream")
    objs[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(paginas)} >>".encode()
    saida, offs = bytearray(b"%PDF-1.4\n"), []
    for i, o in enumerate(objs, 1):
        offs.append(len(saida))
        saida += f"{i} 0 obj\n".encode() + o + b"\nendobj\n"
    xref = len(saida)
    saida += f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode()
    for o in offs:
        saida += f"{o:010d} 00000 n \n".encode()
    saida += f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    caminho.write_bytes(bytes(saida))


def test_pagina_de_imagem_com_carimbo_de_texto_nao_vira_ok(tmp_path):
    carimbo = "Documento assinado digitalmente por FULANO DE TAL em 10/03/2026"
    so_imagem = tmp_path / "digitalizado.pdf"
    _pdf_cru(so_imagem, [("imagem", carimbo)])
    doc, _ = extract.extrair(so_imagem)
    assert doc["status"] == "no_text"
    av = [a for a in doc["warnings"] if a["codigo"] == "pagina_possivelmente_digitalizada"]
    assert len(av) == 1 and av[0]["pagina"] == 1
    corpo = ("Este é o corpo real da petição, com texto de verdade na camada de texto. " * 6)[:300]
    misto = tmp_path / "misto.pdf"
    _pdf_cru(misto, [("imagem", carimbo), ("texto", corpo.replace("(", "").replace(")", ""))])
    doc, _ = extract.extrair(misto)
    assert doc["status"] == "partial" and [a["pagina"] for a in doc["warnings"]
                                           if a["codigo"] == "pagina_possivelmente_digitalizada"] == [1]
    assert any("corpo real" in b["text"] for b in doc["blocks"])


def test_pdf_textual_do_stj_continua_ok_e_sem_aviso_de_digitalizacao():
    doc, _ = extract.extrair(PDF_STJ)
    assert doc["status"] == "ok"
    assert not [a for a in doc["warnings"] if a["codigo"] == "pagina_possivelmente_digitalizada"]
