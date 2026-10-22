"""Rastreabilidade mecânica entre alterações DOCX e apontamentos de revisão."""
import hashlib
import json
import sys
from pathlib import Path

import docx
import pytest

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import extract  # noqa: E402
import redline_docx  # noqa: E402

CLAUSULA = "O prazo é de 12 meses. A multa é de 10%."


@pytest.fixture
def arquivos(tmp_path):
    original, issues, saida = (tmp_path / n for n in ("contrato.docx", "issues.json", "redline.docx"))
    d = docx.Document()
    d.add_paragraph(CLAUSULA)
    d.add_paragraph("O foro é Maceió.")
    d.save(original)
    issues.write_text(json.dumps({"issues": [{"id": "ISS-1", "clause_text": CLAUSULA}]}), encoding="utf-8")
    return original, issues, saida


@pytest.mark.parametrize("edicao, esperado", [
    ({"acao": "substituir", "de": "12 meses", "para": "24 meses"}, "24 meses"),
    ({"acao": "excluir", "de": " A multa é de 10%."}, "O prazo é de 12 meses."),
    ({"acao": "inserir_apos", "ancora": "12 meses.", "texto": " Sem renovação automática."}, "Sem renovação automática."),
])
def test_aplica_edicao_vinculada_e_registra_hash(arquivos, edicao, esperado):
    original, issues, saida = arquivos
    antes = original.read_bytes()
    r = redline_docx.aplicar(original, [{**edicao, "issue_id": "ISS-1"}], saida, issues=issues)
    assert r["aplicadas"] == 1 and r["recusadas"] == 0
    assert r["issues"] == str(issues)
    assert r["issues_sha256"] == hashlib.sha256(issues.read_bytes()).hexdigest()
    assert r["edicoes"][0]["issue_id"] == "ISS-1"
    assert r["edicoes"][0]["vinculo_conferido"] is True
    doc, _ = extract.extrair(saida)
    assert esperado in " ".join(b["text"] for b in doc["blocks"])
    assert original.read_bytes() == antes


@pytest.mark.parametrize("edicao, codigo", [
    ({"acao": "excluir", "de": "12 meses"}, "issue_id_ausente"),
    ({"acao": "excluir", "de": "12 meses", "issue_id": "ISS-9"}, "issue_desconhecida"),
    ({"acao": "excluir", "de": "12 meses", "issue_id": ["ISS-1"]}, "issue_id_ausente"),
    ({"acao": "substituir", "de": "Maceió", "para": "Recife", "issue_id": "ISS-1"}, "trecho_fora_da_issue"),
    ({"acao": "inserir_apos", "ancora": "Maceió", "texto": "!", "issue_id": "ISS-1"}, "trecho_fora_da_issue"),
    ({"acao": "excluir", "de": "", "issue_id": "ISS-1"}, "trecho_fora_da_issue"),
])
def test_recusa_vinculo_invalido_sem_aplicar_edicao(arquivos, edicao, codigo):
    original, issues, saida = arquivos
    r = redline_docx.aplicar(original, [edicao], saida, issues=issues)
    assert r["aplicadas"] == 0 and r["recusadas"] == 1
    assert r["edicoes"][0]["codigo"] == codigo
    assert not r["edicoes"][0]["vinculo_conferido"]
    assert "\n".join(p.text for p in docx.Document(saida).paragraphs) == CLAUSULA + "\nO foro é Maceió."


def test_issue_id_sem_arquivo_de_issues_nao_vira_vinculo_aparente(arquivos):
    original, _, saida = arquivos
    r = redline_docx.aplicar(original, [{"acao": "excluir", "de": "12 meses", "issue_id": "ISS-1"}], saida)
    assert r["aplicadas"] == 0
    assert r["edicoes"][0]["codigo"] == "issues_ausentes"


@pytest.mark.parametrize("obj", [
    [], {}, {"issues": "ISS-1"}, {"issues": [None]},
    {"issues": [{"id": "ISS-1", "clause_text": CLAUSULA}] * 2},
    {"issues": [{"id": "", "clause_text": CLAUSULA}]},
    {"issues": [{"id": "ISS-1", "clause_text": ""}]},
])
def test_arquivo_issues_invalido_preserva_saida_existente(arquivos, obj):
    original, issues, saida = arquivos
    issues.write_text(json.dumps(obj), encoding="utf-8")
    saida.write_bytes(b"saida anterior")
    with pytest.raises(ValueError):
        redline_docx.aplicar(original, [], saida, issues=issues)
    assert saida.read_bytes() == b"saida anterior"


def test_cli_vincula_issues_e_preserva_recusas_parciais(arquivos, tmp_path, capsys):
    original, issues, saida = arquivos
    edicoes = tmp_path / "edicoes.json"
    edicoes.write_text(json.dumps({"edicoes": [
        {"acao": "substituir", "de": "12 meses", "para": "24 meses", "issue_id": "ISS-1"},
        {"acao": "excluir", "de": "Maceió", "issue_id": "ISS-9"},
    ]}), encoding="utf-8")
    assert redline_docx.main([str(original), str(edicoes), str(saida), "--issues", str(issues)]) == 2
    r = json.loads(capsys.readouterr().out)
    assert (r["aplicadas"], r["recusadas"]) == (1, 1)
    assert r["edicoes"][0]["issue_id"] == "ISS-1"
    assert r["edicoes"][1]["codigo"] == "issue_desconhecida"


def test_saida_nao_pode_sobrescrever_issues(arquivos):
    original, issues, _ = arquivos
    antes = issues.read_bytes()
    with pytest.raises(ValueError):
        redline_docx.aplicar(original, [], issues, issues=issues)
    assert issues.read_bytes() == antes


@pytest.mark.parametrize("fonte", ["original", "issues"])
def test_saida_por_hard_link_nao_pode_sobrescrever_entrada(arquivos, fonte):
    original, issues, saida = arquivos
    entrada = original if fonte == "original" else issues
    saida.hardlink_to(entrada)
    antes = entrada.read_bytes()
    with pytest.raises(ValueError):
        redline_docx.aplicar(original, [], saida, issues=issues)
    assert entrada.read_bytes() == antes
    assert saida.read_bytes() == antes


def test_duas_edicoes_podem_remeter_ao_mesmo_apontamento(arquivos):
    original, issues, saida = arquivos
    r = redline_docx.aplicar(original, [
        {"acao": "substituir", "de": "12 meses", "para": "24 meses", "issue_id": "ISS-1"},
        {"acao": "excluir", "de": " A multa é de 10%.", "issue_id": "ISS-1"},
    ], saida, issues=issues)
    assert (r["aplicadas"], r["recusadas"]) == (2, 0)
    assert all(e["issue_id"] == "ISS-1" and e["vinculo_conferido"] for e in r["edicoes"])


def test_vinculo_nao_dispensa_localizacao_no_docx(arquivos):
    original, issues, saida = arquivos
    issues.write_text(json.dumps({"issues": [{"id": "ISS-1", "clause_text": "Cláusula de outro documento."}]}), encoding="utf-8")
    r = redline_docx.aplicar(original, [
        {"acao": "excluir", "de": "outro documento", "issue_id": "ISS-1"},
    ], saida, issues=issues)
    assert r["aplicadas"] == 0 and r["recusadas"] == 1
    assert r["edicoes"][0]["vinculo_conferido"] is True
    assert r["edicoes"][0]["codigo"] == "nao_localizado"


def test_cli_issues_invalido_devolve_erro_json_sem_criar_saida(arquivos, tmp_path, capsys):
    original, issues, saida = arquivos
    issues.write_text("{", encoding="utf-8")
    edicoes = tmp_path / "edicoes.json"
    edicoes.write_text('{"edicoes": []}', encoding="utf-8")
    assert redline_docx.main([str(original), str(edicoes), str(saida), "--issues", str(issues)]) == 1
    assert "erro" in json.loads(capsys.readouterr().out)
    assert not saida.exists()
