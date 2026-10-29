"""redline_docx (controle de alterações) e preencher (modelo com campos de procedência)."""
import hashlib
import json
import sys
import zipfile
from pathlib import Path

import docx
import pytest

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ / "scripts"))
import extract  # noqa: E402
import preencher  # noqa: E402
import redline_docx  # noqa: E402


def sha(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def xml_do_corpo(p):
    with zipfile.ZipFile(p) as z:
        return z.read("word/document.xml").decode("utf-8")


@pytest.fixture
def contrato(tmp_path):
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("Cláusula 5. Em caso de atraso, a multa será de ")
    p.add_run("10% (dez por cento)").bold = True
    p.add_run(" do valor do contrato, sem prejuízo de perdas e danos.")
    d.add_paragraph("Cláusula 6. O prazo é de 60 meses. O foro será o de Maceió.")
    d.add_paragraph("Cláusula 7. Prazo de pagamento: 30 dias. Prazo de entrega: 30 dias.")
    f = tmp_path / "contrato.docx"
    d.save(f)
    return f


# ------------------------------------------------------------------ redline
def test_substituir_trecho_que_atravessa_runs_vira_del_e_ins_e_o_original_fica_intacto(contrato, tmp_path):
    antes = sha(contrato)
    saida = tmp_path / "out" / "revisado.docx"
    r = redline_docx.aplicar(contrato, [{"acao": "substituir", "de": "multa será de 10% (dez por cento) do valor",
                                         "para": "multa será de 2% (dois por cento) do valor"}], saida, autor="Rev")
    assert r["aplicadas"] == 1 and r["recusadas"] == 0 and sha(contrato) == antes == r["original_sha256"]
    xml = xml_do_corpo(saida)
    assert xml.count("<w:del ") == 1 and xml.count("<w:ins ") == 1 and "<w:delText" in xml and 'w:author="Rev"' in xml
    assert "dez por cento" in xml and "dois por cento" in xml
    doc, _ = extract.extrair(saida)
    assert "alteracoes_rastreadas" in [a["codigo"] for a in doc["warnings"]]
    texto = " ".join(b["text"] for b in doc["blocks"])
    assert "dois por cento" in texto and "dez por cento" not in texto  # texto vigente = com a alteração aceita


def test_excluir_e_inserir_apos(contrato, tmp_path):
    saida = tmp_path / "r.docx"
    r = redline_docx.aplicar(contrato, [{"acao": "excluir", "de": ", sem prejuízo de perdas e danos"},
                                        {"acao": "inserir_apos", "ancora": "O foro será o de Maceió", "texto": ", capital do Estado"}], saida)
    assert r["aplicadas"] == 2
    xml = xml_do_corpo(saida)
    assert xml.count("<w:del ") == 1 and xml.count("<w:ins ") == 1
    reaberto = docx.Document(str(saida))
    assert reaberto.paragraphs[0].runs  # continua abrindo no python-docx


def test_edicao_ambigua_ou_inexistente_e_recusada_sem_chute(contrato, tmp_path):
    r = redline_docx.aplicar(contrato, [
        {"acao": "substituir", "de": "30 dias", "para": "15 dias"},          # aparece duas vezes no documento
        {"acao": "substituir", "de": "texto que não existe", "para": "x"},
        {"acao": "voar", "de": "O prazo"},
        {"acao": "substituir", "de": "60 meses", "para": "24 meses"}], tmp_path / "r.docx")
    cods = [x.get("codigo") for x in r["edicoes"]]
    assert cods == ["ambiguo", "nao_localizado", "acao_invalida", None]
    assert r["aplicadas"] == 1 and r["recusadas"] == 3
    xml = xml_do_corpo(tmp_path / "r.docx")
    assert xml.count("<w:del ") == 1 and "24 meses" in xml


def test_trecho_que_cruza_hiperlink_ou_tabulacao_e_recusado(tmp_path):
    d = docx.Document()
    p = d.add_paragraph("Veja o item ")
    p.add_run("A")
    p.add_run().add_tab()
    p.add_run("B final.")
    f = tmp_path / "t.docx"
    d.save(f)
    r = redline_docx.aplicar(f, [{"acao": "excluir", "de": "item AB final"}], tmp_path / "o.docx")
    assert r["edicoes"][0]["resultado"] == "recusada" and r["edicoes"][0]["codigo"] in ("nao_localizado", "trecho_cruza_estrutura")


def test_saida_igual_ao_original_e_proibida(contrato):
    with pytest.raises(ValueError, match="próprio original"):
        redline_docx.aplicar(contrato, [], contrato)


def test_cli_devolve_codigo_2_quando_ha_recusa(contrato, tmp_path):
    ed = tmp_path / "e.json"
    ed.write_text(json.dumps({"edicoes": [{"acao": "excluir", "de": "não existe"}]}), encoding="utf-8")
    assert redline_docx.main([str(contrato), str(ed), str(tmp_path / "o.docx")]) == 2


# ------------------------------------------------------------------ preencher
@pytest.fixture
def modelo(tmp_path):
    d = docx.Document()
    p = d.add_paragraph()
    p.add_run("CONTRATANTE: ")
    p.add_run("{{contratante").bold = True
    p.add_run("}}, CNPJ {{cnpj}}.")  # marcador dividido em dois runs
    d.add_paragraph("Prazo: {{prazo}}. Foro: {{foro}}. Multa: {{multa}}. Reajuste: {{reajuste}}. Garantia: {{garantia}}.")
    f = tmp_path / "modelo.docx"
    d.save(f)
    return f


def campos_base():
    return {"campos": [
        {"nome": "contratante", "estado": "extraido", "valor": "Alfa Serviços Ltda.",
         "fonte": {"document_id": "doc-1", "locator": "par.2", "trecho": "A contratante é Alfa Serviços Ltda., com sede em Maceió."}},
        {"nome": "cnpj", "estado": "pendente", "motivo": "o CNPJ não consta dos documentos"},
        {"nome": "prazo", "estado": "informado", "valor": "24 meses", "informado_em": "2026-10-17", "fala": "O prazo será de 24 meses."},
        {"nome": "foro", "estado": "inferido", "valor": "Maceió", "razao": "sede da contratante"},
        {"nome": "multa", "estado": "conflito", "candidatos": [{"valor": "2%"}, {"valor": "10%"}]},
        {"nome": "reajuste", "estado": "nao_se_aplica", "motivo": "contrato de prazo certo sem reajuste"},
        {"nome": "garantia", "estado": "pendente", "motivo": "a definir"}]}


def texto_docx(p):
    return "\n".join(x.text for x in docx.Document(str(p)).paragraphs)


def test_gera_preenchendo_so_o_que_tem_prova_e_marcando_o_resto(modelo, tmp_path):
    antes = sha(modelo)
    r = preencher.gerar(campos_base(), modelo, tmp_path / "o.docx")
    assert r["gerado"] and not r["completo"] and sha(modelo) == antes
    t = texto_docx(tmp_path / "o.docx")
    assert "CONTRATANTE: Alfa Serviços Ltda., CNPJ [PENDENTE: cnpj]." in t
    assert "Prazo: 24 meses." in t and "Foro: [PENDENTE (inferido não aceito): foro]" in t
    assert "[CONFLITO: multa]" in t and "[NÃO SE APLICA: reajuste]" in t and "{{" not in t
    assert set(r["marcados_no_documento"]) == {"cnpj", "foro", "multa", "reajuste", "garantia"}
    assert set(r["preenchidos"]) == {"contratante", "prazo"}
    assert 'w:highlight w:val="yellow"' in xml_do_corpo(tmp_path / "o.docx")


def test_inferido_so_entra_com_aceite_explicito_e_marcado(modelo, tmp_path):
    r = preencher.gerar(campos_base(), modelo, tmp_path / "o.docx", aceitar_inferidos=True)
    assert "Foro: Maceió [INFERIDO: confirmar]" in texto_docx(tmp_path / "o.docx") and "foro" not in r["marcados_no_documento"]


def test_completo_so_quando_nada_ficou_marcado(modelo, tmp_path):
    todos = {"campos": [{"nome": n, "estado": "informado", "valor": f"valor {n}", "informado_em": "2026-10-17", "fala": f"o campo é valor {n}"}
                        for n in ("contratante", "cnpj", "prazo", "foro", "multa", "reajuste", "garantia")]}
    r = preencher.gerar(todos, modelo, tmp_path / "o.docx")
    assert r["gerado"] and r["completo"] is True and r["marcados_no_documento"] == {}
    t = texto_docx(tmp_path / "o.docx")
    assert "CONTRATANTE: valor contratante, CNPJ valor cnpj." in t and "{{" not in t and "[" not in t


@pytest.mark.parametrize("mutar, codigo", [
    (lambda c: c[0]["fonte"].update(trecho="Texto qualquer sem o valor."), "valor_fora_do_trecho"),
    (lambda c: c[0].pop("fonte"), "fonte_incompleta"),
    (lambda c: c[1].update(valor="12.345.678/0001-90"), "valor_em_estado_sem_valor"),
    (lambda c: c[1].pop("motivo"), "motivo_ausente"),
    (lambda c: c[2].pop("informado_em"), "informado_em_invalido"),
    (lambda c: c[2].pop("fala"), "informado_sem_fala"),
    (lambda c: c[2].update(fala="Ponha o mais provável."), "valor_fora_da_fala"),
    (lambda c: c.append({"nome": "poderes", "estado": "inferido", "valor": "amplos", "razao": "mais provável"}), "inferido_em_campo_sensivel"),
    (lambda c: c[3].pop("razao"), "razao_ausente"),
    (lambda c: c[4].update(candidatos=[{"valor": "2%"}, {"valor": "2%"}]), "candidatos_insuficientes"),
    (lambda c: c[0].update(estado="certeza"), "estado_invalido"),
    (lambda c: c[1].update(nome="cnpj"), None),
    (lambda c: c.append({"nome": "prazo", "estado": "pendente", "motivo": "x"}), "nome_duplicado"),
    (lambda c: c.append({"nome": "Nome Ruim", "estado": "pendente", "motivo": "x"}), "nome_invalido"),
])
def test_lint_barra_campo_sem_prova(mutar, codigo):
    c = campos_base()
    mutar(c["campos"])
    cods = [e["codigo"] for e in preencher.lint(c)["erros"]]
    if codigo is None:
        assert cods == []
    else:
        assert codigo in cods


def test_lint_cruza_com_o_modelo_e_com_a_extracao(modelo):
    c = campos_base()
    c["campos"].append({"nome": "extra", "estado": "pendente", "motivo": "x"})
    c["campos"] = [x for x in c["campos"] if x["nome"] != "garantia"]
    r = preencher.lint(c, modelo)
    assert [e["codigo"] for e in r["erros"]] == ["marcador_sem_campo"] and r["erros"][0]["onde"] == "garantia"
    assert [a["codigo"] for a in r["avisos"] if a["codigo"] == "campo_sem_uso"] == ["campo_sem_uso"]
    ext = {"documents": [{"document_id": "doc-1", "blocks": [
        {"id": "b1", "locator": "docx.par.2", "text": "A contratante é Alfa Serviços Ltda., com sede em Maceió."}]}]}
    ok = campos_base()
    assert preencher.lint(ok, extracoes=ext)["erros"] == []
    ok["campos"][0]["fonte"]["locator"] = "par.9"
    assert "localizador_inexistente" in [e["codigo"] for e in preencher.lint(ok, extracoes=ext)["erros"]]
    ok["campos"][0]["fonte"].update(locator="par.2", document_id="doc-9")
    assert "documento_desconhecido" in [e["codigo"] for e in preencher.lint(ok, extracoes=ext)["erros"]]
    ok["campos"][0]["fonte"].update(document_id="doc-1", trecho="A contratante é Alfa Serviços Ltda.")
    ext["documents"][0]["blocks"][0]["text"] = "Outra coisa, Alfa Serviços Ltda. aparece."
    assert "trecho_fora_do_bloco" in [e["codigo"] for e in preencher.lint(ok, extracoes=ext)["erros"]]


def test_gerar_recusa_quando_o_lint_tem_erro(modelo, tmp_path):
    c = campos_base()
    c["campos"][0]["fonte"]["trecho"] = "nada a ver"
    r = preencher.gerar(c, modelo, tmp_path / "o.docx")
    assert r["gerado"] is False and not (tmp_path / "o.docx").exists()


def test_trecho_que_e_so_o_valor_gera_aviso_e_o_relatorio_manda_escalar_poderes(modelo, tmp_path):
    c = campos_base()
    c["campos"][0]["fonte"]["trecho"] = "Alfa Serviços Ltda."
    avisos = [a["codigo"] for a in preencher.lint(c)["avisos"]]
    assert "trecho_minimo" in avisos
    d = {"campos": [{"nome": n, "estado": "pendente", "motivo": "sem dado"} for n in ("contratante", "cnpj", "prazo", "foro", "multa", "reajuste", "garantia")]}
    r = preencher.gerar(d, modelo, tmp_path / "o.docx")
    assert r["escalar"] and "terço" in r["escalar"][0]


def test_cpf_diferente_nos_documentos_gera_possivel_conflito():
    ext = {"documents": [
        {"document_id": "doc-1", "blocks": [{"id": "b1", "locator": "docx.par.1", "text": "Outorgante: Maria, CPF 111.444.777-35, residente em Maceió."}]},
        {"document_id": "doc-2", "blocks": [{"id": "b1", "locator": "eml.corpo.par.1", "text": "Meu CPF é 222.555.888-46, favor usar."}]}]}
    c = {"campos": [{"nome": "cpf", "estado": "extraido", "valor": "111.444.777-35",
                     "fonte": {"document_id": "doc-1", "locator": "par.1", "trecho": "Outorgante: Maria, CPF 111.444.777-35, residente em Maceió."}}]}
    r = preencher.lint(c, extracoes=ext)
    assert r["erros"] == [] and "possivel_conflito" in [a["codigo"] for a in r["avisos"]]


def test_localizador_por_sufixo_escolhe_o_bloco_que_contem_o_trecho():
    """Dois blocos terminam em 'par.2': o primeiro não pode ser tomado só por vir antes."""
    trecho = "A contratante é Alfa Serviços Ltda., com sede em Maceió."
    ext = {"documents": [{"document_id": "doc-1", "blocks": [
        {"id": "b0", "locator": "tabela.par.2", "text": "Outro texto, sem relação com a contratante."},
        {"id": "b1", "locator": "docx.par.2", "text": trecho}]}]}
    c = campos_base()
    assert preencher.lint(c, extracoes=ext)["erros"] == []
    # sem nenhum candidato que contenha o trecho, o erro continua sendo apontado
    ext["documents"][0]["blocks"][1]["text"] = "Texto diferente."
    assert "trecho_fora_do_bloco" in [e["codigo"] for e in preencher.lint(c, extracoes=ext)["erros"]]


def test_localizador_exato_tem_prioridade_sobre_o_sufixo():
    trecho = "A contratante é Alfa Serviços Ltda., com sede em Maceió."
    ext = {"documents": [{"document_id": "doc-1", "blocks": [
        {"id": "b0", "locator": "docx.par.2", "text": "Texto de outro bloco."},
        {"id": "b1", "locator": "par.2", "text": trecho}]}]}
    assert preencher.lint(campos_base(), extracoes=ext)["erros"] == []
